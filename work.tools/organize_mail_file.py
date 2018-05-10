#!/usr/bin/env python3
# -*- coding: utf-8 -*-

__author__ = 'Cavin Cao'

'''
	功能：整理目标文件夹下的文件，生成对应的excel
'''

import xlsxwriter
import subprocess
import os
import docx
import sys
import re
#import pypinyin
#from pypinyin import lazy_pinyin


path='/Users/cavin/Desktop/files'

'''
    文件重命名：{index}-{中文}.docx
'''
def remove_file_space():
    for filename in os.listdir(path):
        otherName = re.sub("[\s+\!\/_,$%^*(+\"\')]+|[+——()?【】“”！，。？、~@#￥%……&*（）]+", "",filename)  
        #otherName=filename.replace(' ','').replace('&','').replace('(','').replace(')','')
        os.rename(os.path.join(path,filename),os.path.join(path,otherName))

def rename_files(filename,targetChineseName,index):
    eName=targetChineseName #''.join(lazy_pinyin(targetChineseName))
    otherName=str(index)+'-'+eName+os.path.splitext(filename)[1]
    os.rename(os.path.join(path,filename),os.path.join(path,otherName))
    #return os.path.join('file://',path,otherName)
    return otherName

def etl_word_files(fullname):
    doc = docx.Document(fullname)
    sheetModel={}
    specialDeal=0
    for para in doc.paragraphs:
        try:
            paraText=para.text.replace('：',':')#统一转成英文，方便split
            if specialDeal==1:
                if paraText.strip()!='':
                    #存在多个空格的问题
                    paraText=' '.join(paraText.split())
                    sheetModel['scollage']=paraText.split(' ')[1].strip()
                    specialDeal=0
            elif specialDeal==2:
                if paraText.strip()!='':
                    paraText=' '.join(paraText.split())
                    sheetModel['swork']=paraText.split(' ')[1].strip()
                    specialDeal=0
            elif '姓名:' in paraText:
                sname=paraText.split(':',1)[1].strip()
                if(len(sname)>3):
                    sname=sname[0:4]
                sheetModel['sname']=paraText.split(':',1)[1].strip()
            elif '性别:' in paraText:
                sheetModel['ssex']=paraText.split(':',1)[1].strip()
            elif '年龄:' in paraText:
                sheetModel['sage']=paraText.split(':',1)[1].strip()
            elif '籍贯:' in paraText:
                sheetModel['sprovince']=paraText.split(':',1)[1].strip()
            elif '目前所在地:' in paraText:
                sheetModel['slocation']=paraText.split(':',1)[1].strip()
            elif '学历:' in paraText:
                sheetModel['seducation']=paraText.split(':',1)[1].strip()
            elif '教育背景' in paraText:
                specialDeal=1#教育背景取下一行数据
            elif ('工作经历' in paraText or '工作经验' in paraText) and ':' in paraText:
                specialDeal=2
            elif '职位名称:' in paraText:
                sheetModel['stitle']=paraText.split(':',1)[1].strip()
            else:
                pass
        except Exception as e:
            print('Error:'+fullname+'****'+paraText)
            print(e)
    return sheetModel
        

def organize_mail_mian():
    workbook = xlsxwriter.Workbook('report_list.xlsx')
    worksheet = workbook.add_worksheet('list')
    worksheet.write(0,0, '序号') 
    worksheet.write(0,1, '姓名') 
    worksheet.write(0,2, '性别') 
    worksheet.write(0,3, '年龄') 
    worksheet.write(0,4, '籍贯') 
    worksheet.write(0,5, '目前所在地') 
    worksheet.write(0,6, '学历')
    worksheet.write(0,7, '学校')
    worksheet.write(0,8, '公司')
    worksheet.write(0,9, '职位')
    worksheet.write(0,10, '文档链接')

    index=0
    for filename in os.listdir(path):
        if filename=='.DS_Store': #操作之后mac会出现Ds_Store文件？
            continue
        index=index+1
        sheetModel={}
        fullname=os.path.join(path,filename)
        if filename.endswith('.doc'):
            subprocess.call('textutil -convert docx {0}'.format(fullname),shell=True)
            fullname=fullname[:-4]+".docx"
            sheetModel= etl_word_files(fullname)
            subprocess.call('rm {0}'.format(fullname),shell=True) #移除转换的文件
        elif filename.endswith('.docx'):
            sheetModel= etl_word_files(fullname)
        else:
            sheetModel={}

        worksheet.write(index,0, index) 
        worksheet.write(index,1, sheetModel.get('sname','')) 
        worksheet.write(index,2, sheetModel.get('ssex','')) 
        worksheet.write(index,3, sheetModel.get('sage','')) 
        worksheet.write(index,4, sheetModel.get('sprovince','')) 
        worksheet.write(index,5, sheetModel.get('slocation','')) 
        worksheet.write(index,6, sheetModel.get('seducation',''))
        worksheet.write(index,7, sheetModel.get('scollage',''))
        worksheet.write(index,8, sheetModel.get('swork',''))
        worksheet.write(index,9, sheetModel.get('stitle',''))
        #otherName=rename_files(filename,sheetModel.get('sname',''),index)      
        worksheet.write(index,10, '=HYPERLINK(\"./'+filename+'\",\"附件\")')
        #worksheet.write_url(index,10, otherName, string='附件') 
        print(str(index)+':完成')
    

    workbook.close()

if __name__ =='__main__':
    remove_file_space()
    organize_mail_mian()

        
            